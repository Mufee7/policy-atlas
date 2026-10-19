"""Rebuild fictional ingestion fixtures (developer-only: reportlab, Pillow, python-docx)."""
from pathlib import Path
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.lib.pdfencrypt import StandardEncryption
from PIL import Image,ImageDraw
from docx import Document
root=Path(__file__).resolve().parents[1]/'public'/'samples'
root.mkdir(parents=True,exist_ok=True)
def page(c,title,lines):
    c.setFont('Helvetica-Bold',16);c.drawString(60,780,title)
    c.setFont('Helvetica',11)
    for i,line in enumerate(lines):c.drawString(60,745-i*19,line)
    c.showPage()
c=canvas.Canvas(str(root/'travel-policy.pdf'))
page(c,'Fictional Travel Policy - Version 1.0',['SYNTHETIC DEMO DOCUMENT - not real company guidance.','1. Travel Approval','Business travel requires manager approval before booking.'])
page(c,'2. Per Diem',['The domestic meal allowance is LKR 3,500 per travel day.','Receipts are required. Alcohol is not reimbursable.'])
c.save()
im=Image.new('RGB',(950,380),'white');d=ImageDraw.Draw(im);d.text((40,60),'FICTIONAL IMAGE-ONLY POLICY PAGE',fill='black');d.text((40,110),'This text is an image. It requires OCR to become searchable.',fill='black')
for name,with_text in [('mixed-scan.pdf',True),('image-only.pdf',False)]:
    c=canvas.Canvas(str(root/name))
    if with_text:page(c,'Fictional Equipment Policy',['SYNTHETIC DEMO DOCUMENT.','Employees must return company equipment within five days of departure.'])
    c.drawImage(ImageReader(im),50,350,width=500,height=200);c.showPage();c.save()
c=canvas.Canvas(str(root/'protected.pdf'),encrypt=StandardEncryption('demo-password'))
page(c,'Protected fictional document',['This sample must be rejected without asking for a password.']);c.save()
doc=Document();doc.add_heading('Fictional Access Request Procedure',0);doc.add_paragraph('SYNTHETIC DEMO DOCUMENT - not real company guidance.');doc.add_heading('1. Access Requests',1);doc.add_paragraph('Employees must submit access requests through the service desk portal.');doc.add_paragraph('The system owner must approve access before provisioning.');doc.add_heading('2. Removal',1);doc.add_paragraph('Report access removal requests to the service desk within one business day.');doc.save(root/'access-policy.docx')
print('Generated fictional PDF and DOCX fixtures.')
